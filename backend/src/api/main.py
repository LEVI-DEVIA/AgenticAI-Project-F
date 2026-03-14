from __future__ import annotations

from dotenv import find_dotenv, load_dotenv

# Charge les variables d'environnement en toute première chose, en cherchant activement le .env
load_dotenv(find_dotenv())

import asyncio
import csv
import json
import os
import re
from pathlib import Path
from time import time
from uuid import uuid4

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from google import genai
from google.genai import types
from openpyxl import load_workbook
from src.agent.agent_avi import (
    DEFAULT_LINK,
    create_browserbase_session,
    get_browserbase_live_view_urls,
    run_agent_stagehand,
)
from sse_starlette.sse import EventSourceResponse

app = FastAPI(
    title="AgentAI-Project-F",
    version="0.0.1",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


RUNS: dict[str, dict] = {}


def _alloc_run_id(filename: str) -> str:
    return uuid4().hex


def _xlsx_to_csv_bytes(xlsx_path: Path) -> bytes:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel vide")
    header = rows[0]
    if not header:
        raise ValueError("Excel sans en-têtes")

    out_lines: list[list[str]] = []
    out_lines.append(["" if v is None else str(v) for v in header])
    for r in rows[1:]:
        if r is None:
            continue
        out_lines.append(["" if v is None else str(v) for v in r])

    from io import StringIO

    sio = StringIO()
    w = csv.writer(sio)
    w.writerows(out_lines)
    return sio.getvalue().encode("utf-8")


def _extract_json_object(text: str) -> dict:
    m = re.search(r"\{.*\}", text, re.DOTALL)
    if not m:
        raise ValueError("Aucun JSON détecté")
    raw = m.group(0)
    return json.loads(raw)


async def _start_run_from_csv_path(csv_path: Path) -> dict:
    if not DEFAULT_LINK:
        raise HTTPException(status_code=500, detail="DEFAULT_LINK manquant")

    run_id = _alloc_run_id(csv_path.name)

    try:
        session = await create_browserbase_session()
        live_urls = await get_browserbase_live_view_urls(session["id"])
    except RuntimeError as e:
        raise HTTPException(status_code=500, detail=str(e))

    submit_event = asyncio.Event()
    events: asyncio.Queue[dict] = asyncio.Queue()

    RUNS[run_id] = {
        "run_id": run_id,
        "state": "running",
        "created_at": time(),
        "csv_path": str(csv_path),
        "browserbase": {
            "session_id": session["id"],
            "connect_url": session.get("connectUrl"),
            "selenium_remote_url": session.get("seleniumRemoteUrl"),
        },
        "live_view": {
            "debugger_url": live_urls.get("debuggerUrl"),
            "debugger_fullscreen_url": live_urls.get("debuggerFullscreenUrl"),
        },
        "current_row": None,
        "_submit_event": submit_event,
        "_events": events,
    }

    async def runner():
        try:
            run_state = RUNS.get(run_id)
            if not run_state:
                return
            session_id = (run_state.get("browserbase") or {}).get("session_id")
            if not isinstance(session_id, str):
                raise RuntimeError("session_id manquant ou invalide")

            async def wait_for_user_submit(row_index: int):
                run_state_inner = RUNS.get(run_id)
                if not run_state_inner:
                    return
                run_state_inner["current_row"] = row_index
                run_state_inner["state"] = "awaiting_user_submit"
                q: asyncio.Queue = run_state_inner["_events"]
                await q.put({"type": "awaiting_user_submit", "row_index": row_index})

                event: asyncio.Event = run_state_inner["_submit_event"]
                await event.wait()
                event.clear()
                run_state_inner["state"] = "running"
                await q.put({"type": "running", "row_index": row_index})

            await run_agent_stagehand(
                Path(run_state["csv_path"]),
                session_id=session_id,
                wait_for_user_submit=wait_for_user_submit,
            )

            run_state["state"] = "completed"
            q_done: asyncio.Queue = run_state["_events"]
            await q_done.put({"type": "completed"})
        except Exception as e:
            run_state = RUNS.get(run_id)
            if run_state:
                run_state["state"] = "failed"
                q_err: asyncio.Queue = run_state.get("_events")
                if isinstance(q_err, asyncio.Queue):
                    await q_err.put({"type": "failed", "error": str(e)})

    asyncio.create_task(runner())
    return {
        "run_id": run_id,
        "live_view_url": RUNS[run_id]["live_view"]["debugger_fullscreen_url"],
    }


@app.post("/run")
async def run(file: UploadFile = File(...)):
    filename = (file.filename or "").lower()
    is_csv = filename.endswith(".csv")
    is_xlsx = filename.endswith(".xlsx")

    if not (is_csv or is_xlsx):
        raise HTTPException(
            status_code=400, detail="Le fichier doit être un CSV ou un Excel (.xlsx)"
        )

    backend_dir = Path(__file__).resolve().parents[2]
    upload_dir = backend_dir / "data" / "uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)

    run_id = _alloc_run_id(file.filename or "upload")
    raw_path = upload_dir / f"{uuid4()}__{file.filename or 'upload'}"
    content = await file.read()
    raw_path.write_bytes(content)

    if is_xlsx:
        try:
            csv_bytes = _xlsx_to_csv_bytes(raw_path)
        except Exception as e:
            raise HTTPException(status_code=400, detail=str(e))
        out_path = upload_dir / f"{uuid4()}__{run_id}.csv"
        out_path.write_bytes(csv_bytes)
    else:
        out_path = upload_dir / f"{uuid4()}__{run_id}.csv"
        out_path.write_bytes(content)

    return await _start_run_from_csv_path(out_path)


@app.post("/run/image")
async def run_from_image(file: UploadFile = File(...)):
    content_type = file.content_type or ""
    if not content_type.startswith("image/"):
        raise HTTPException(status_code=400, detail="Le fichier doit être une image")

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY manquant")

    image_bytes = await file.read()
    if not image_bytes:
        raise HTTPException(status_code=400, detail="Image vide")

    client = genai.Client(api_key=api_key)
    prompt = (
        "Tu fais de l'OCR et extraction de champs. "
        "Retourne uniquement un objet JSON (sans texte autour). "
        "Si tu ne connais pas une valeur, mets null. "
        "Essaie d'extraire des champs utiles (nom, prénom, email, téléphone, adresse, âge, etc.) en fonction de ce que tu vois."
    )

    resp = client.models.generate_content(
        model="gemini-3-flash-preview",
        contents=[
            types.Part.from_bytes(data=image_bytes, mime_type=content_type),
            prompt,
        ],
    )

    try:
        extracted = _extract_json_object(resp.text or "")
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"OCR JSON invalide: {e}. Texte: {resp.text}",
        )

    if not isinstance(extracted, dict) or not extracted:
        raise HTTPException(status_code=422, detail="OCR vide ou non exploitable")

    backend_dir = Path(__file__).resolve().parents[2]
    upload_dir = backend_dir / "data" / "uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)

    run_id = _alloc_run_id(file.filename or "image")
    out_path = upload_dir / f"{uuid4()}__{run_id}.csv"

    from io import StringIO

    sio = StringIO()
    writer = csv.DictWriter(sio, fieldnames=list(extracted.keys()))
    writer.writeheader()
    writer.writerow(extracted)
    out_path.write_text(sio.getvalue(), encoding="utf-8")

    return await _start_run_from_csv_path(out_path)


@app.post("/ocr")
async def ocr_image(file: UploadFile = File(...)):
    content_type = file.content_type or ""
    if not content_type.startswith("image/"):
        raise HTTPException(status_code=400, detail="Le fichier doit être une image")

    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY manquant")

    image_bytes = await file.read()
    if not image_bytes:
        raise HTTPException(status_code=400, detail="Image vide")

    client = genai.Client(api_key=api_key)
    prompt = (
        "Tu fais de l'OCR et extraction de champs. "
        "Retourne uniquement un objet JSON (sans texte autour). "
        "Si tu ne connais pas une valeur, mets null. "
        "Essaie d'extraire des champs utiles (nom, prénom, email, téléphone, adresse, âge, etc.) en fonction de ce que tu vois."
    )

    resp = client.models.generate_content(
        model="gemini-3-flash-preview",
        contents=[
            types.Part.from_bytes(data=image_bytes, mime_type=content_type),
            prompt,
        ],
    )

    try:
        data = _extract_json_object(resp.text or "")
    except Exception as e:
        raise HTTPException(
            status_code=422, detail=f"OCR JSON invalide: {e}. Texte: {resp.text}"
        )

    return {"data": data, "raw_text": resp.text}


@app.get("/run/{run_id}/events")
async def run_events(run_id: str):
    run_state = RUNS.get(run_id)
    if not run_state:
        raise HTTPException(status_code=404, detail="Run introuvable")

    async def event_generator():
        q: asyncio.Queue = run_state["_events"]
        while True:
            evt = await q.get()
            yield {
                "event": "message",
                "data": json.dumps(evt, ensure_ascii=False),
            }
            if evt.get("type") in {"completed", "failed"}:
                break

    return EventSourceResponse(event_generator())


@app.post("/run/{run_id}/submit")
async def submit_row(run_id: str):
    run_state = RUNS.get(run_id)
    if not run_state:
        raise HTTPException(status_code=404, detail="Run introuvable")

    event = run_state.get("_submit_event")
    if isinstance(event, asyncio.Event):
        event.set()
    return {"ok": True}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
